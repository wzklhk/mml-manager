from openpyxl import Workbook, load_workbook

from app.converters.mml_to_xls import convert_file_to_excel_and_csv
from app.utils.table import MmlConfig, MmlDataSet
from app.utils.tabular import read_csv, read_excel, write_excel


def test_write_excel_preserves_numeric_and_string_cell_types(tmp_path):
    dataset = MmlDataSet()
    dataset.add(
        MmlConfig(
            table="TYPES",
            values={"INT": 7, "FLOAT": 1.25, "TEXT": "007", "FORMULA_TEXT": "=1+1"},
        )
    )
    output = tmp_path / "types.xlsx"

    write_excel(str(output), dataset)

    workbook = load_workbook(output, read_only=True, data_only=True)
    sheet = workbook["TYPES"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, headers["INT"]).value == 7
    assert sheet.cell(2, headers["INT"]).data_type == "n"
    assert sheet.cell(2, headers["FLOAT"]).value == 1.25
    assert sheet.cell(2, headers["FLOAT"]).data_type == "n"
    assert sheet.cell(2, headers["TEXT"]).value == "007"
    assert sheet.cell(2, headers["TEXT"]).data_type == "s"
    assert sheet.cell(2, headers["TEXT"]).number_format == "@"
    assert sheet.cell(2, headers["FORMULA_TEXT"]).value == "=1+1"
    assert sheet.cell(2, headers["FORMULA_TEXT"]).data_type == "s"
    workbook.close()


def test_legacy_excel_converter_preserves_value_types(tmp_path):
    input_path = tmp_path / "types.mml"
    input_path.write_text('SET TYPES:NUMBER=42,TEXT="042";', encoding="utf-8")

    result = convert_file_to_excel_and_csv(str(input_path), str(tmp_path / "converted"))

    workbook = load_workbook(result["excel_path"], read_only=True, data_only=True)
    sheet = workbook["TYPES"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, headers["NUMBER"]).value == 42
    assert sheet.cell(2, headers["NUMBER"]).data_type == "n"
    assert sheet.cell(2, headers["TEXT"]).value == "042"
    assert sheet.cell(2, headers["TEXT"]).data_type == "s"
    assert sheet.cell(2, headers["TEXT"]).number_format == "@"
    workbook.close()


def test_tabular_readers_preserve_and_infer_scalar_types(tmp_path):
    excel_path = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TYPES"
    sheet.append(["NUMBER", "CODE"])
    sheet.append([42, "042"])
    workbook.save(excel_path)
    workbook.close()

    csv_path = tmp_path / "input.csv"
    csv_path.write_text("NUMBER,CODE,ENABLED\n42,042,true\n", encoding="utf-8")

    excel_values = read_excel(str(excel_path)).get_group("TYPES").configs[0].values
    csv_values = read_csv(str(csv_path)).get_group("input").configs[0].values
    assert excel_values == {"NUMBER": 42, "CODE": "042"}
    assert csv_values == {"NUMBER": 42, "CODE": "042", "ENABLED": True}
