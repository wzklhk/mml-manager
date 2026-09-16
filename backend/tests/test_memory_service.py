import csv
import io
import zipfile

import pytest
from openpyxl import load_workbook

from app.repositories import sqlite as repository
from app.services import mml as mml_service


def test_import_view_edit_and_export_use_memory(tmp_path):
    source = tmp_path / "config.txt"
    source.write_text(
        "ADD USER PROFILE:ID=2,NAME='Jane Doe';\n" "SET USER PROFILE:ID=1,NAME=Admin;\n",
        encoding="utf-8",
    )

    imported = mml_service.import_mml_file(str(source))
    summary = mml_service.get_tables_summary()
    page = mml_service.get_configs("USER PROFILE", page=1, page_size=20)

    assert imported["total_count"] == 2
    assert summary[0]["count"] == 2
    assert [row["cmd_type"] for row in page["configs"]] == ["SET", "ADD"]

    row_id = page["configs"][0]["id"]
    assert mml_service.update_config("USER PROFILE", row_id, {"ID": "1", "NAME": "Root User"})
    exported = mml_service.export_mml("USER PROFILE")["content"]
    assert 'SET USER PROFILE:ID=1,NAME="Root User";' in exported
    assert "ADD USER PROFILE:ID=2,NAME=Jane Doe;" not in exported
    assert 'ADD USER PROFILE:ID=2,NAME="Jane Doe";' in exported


def test_each_import_is_retained_and_can_be_activated(tmp_path):
    first = tmp_path / "first.mml"
    second = tmp_path / "second.mml"
    first.write_text("SET FIRST:ID=1;", encoding="utf-8")
    second.write_text("SET SECOND:ID=2;", encoding="utf-8")

    first_result = mml_service.import_mml_file(str(first))
    second_result = mml_service.import_mml_file(str(second))

    assert [table["table_name"] for table in mml_service.get_tables_summary()] == ["SECOND"]
    snapshots = mml_service.get_snapshots()
    assert len(snapshots["snapshots"]) >= 2
    assert snapshots["active_id"] == second_result["snapshot"]["id"]

    mml_service.activate_snapshot(first_result["snapshot"]["id"])
    assert [table["table_name"] for table in mml_service.get_tables_summary()] == ["FIRST"]


def test_empty_configuration_can_be_created_and_populated():
    created = mml_service.create_configuration("Manual Config")

    assert created["total_count"] == 0
    assert created["snapshot"]["name"] == "Manual Config"
    assert mml_service.get_snapshots()["active_id"] == created["snapshot"]["id"]
    assert mml_service.get_tables_summary() == []

    mml_service.create_table("CELL", ["ID"])
    assert [table["table_name"] for table in mml_service.get_tables_summary()] == ["CELL"]


def test_imported_configuration_uses_only_the_original_file_name():
    imported = mml_service.import_configuration_stream(
        io.BytesIO(b"SET CELL:ID=1;"),
        r"C:\\fakepath\\radio-config.mml",
    )

    assert imported["snapshot"]["name"] == "radio-config.mml"


def test_table_summary_is_sorted_by_command_name():
    mml_service.import_mml_text("SET ZEBRA:ID=1; SET alpha:ID=2; SET Middle:ID=3;", "sorted.mml")
    assert [table["table_name"] for table in mml_service.get_tables_summary()] == ["alpha", "Middle", "ZEBRA"]


def test_imported_table_column_types_are_inferred_from_values():
    mml_service.import_mml_text(
        'SET TYPES:COUNT=7,RATIO=1.5,ACTIVE=true,CODE="007"; ' 'SET TYPES:COUNT=8,RATIO=2,ACTIVE=false,CODE="008";',
        "types.mml",
    )

    assert mml_service.get_tables_summary()[0]["column_types"] == {
        "ACTIVE": "boolean",
        "CODE": "string",
        "COUNT": "integer",
        "RATIO": "decimal",
    }


def test_legacy_table_without_type_metadata_is_inferred_and_backfilled():
    imported = mml_service.import_mml_text("SET LEGACY:ID=1,NAME=Alpha;", "legacy.mml")
    with repository.DatabaseConnection() as db:
        db.execute(
            "UPDATE _mml_snapshot_tables SET column_types_json='{}', type_inference_version=0 "
            "WHERE snapshot_id=? AND table_name='LEGACY'",
            (imported["snapshot"]["id"],),
        )
    mml_service.store._cache.invalidate()

    summary = mml_service.get_tables_summary()[0]

    assert summary["column_types"] == {"ID": "integer", "NAME": "string"}
    with repository.DatabaseConnection() as db:
        stored = db.execute(
            "SELECT column_types_json FROM _mml_snapshot_tables WHERE snapshot_id=? AND table_name='LEGACY'",
            (imported["snapshot"]["id"],),
        ).fetchone()["column_types_json"]
    assert '"ID": "integer"' in stored


def test_configs_can_be_filtered_by_multiple_columns():
    mml_service.import_mml_text(
        'SET CELL:ID=1,NAME="Alpha_100%"; SET CELL:ID=2,NAME="AlphaX100Y"; SET CELL:ID=3,NAME="Beta";',
        "filters.mml",
    )

    filtered = mml_service.get_configs("CELL", page_size=20, filters={"NAME": "alpha_100%", "ID": "1"})

    assert filtered["total"] == 1
    assert [row["config_data"]["ID"] for row in filtered["configs"]] == [1]


def test_tables_can_be_created_and_deleted_in_the_active_snapshot():
    mml_service.import_mml_text("SET EXISTING:ID=1;", "tables.mml")

    created = mml_service.create_table("NEW TABLE", ["NAME", "ID", "NAME"], {"ID": "integer", "NAME": "string"})
    row_id = mml_service.add_config("NEW TABLE", {"ID": 7, "NAME": "Alpha"})

    assert created == {
        "table_name": "NEW TABLE",
        "columns": ["ID", "NAME"],
        "column_types": {"ID": "integer", "NAME": "string"},
        "count": 0,
    }
    assert mml_service.get_config("NEW TABLE", row_id)["config_data"] == {"ID": 7, "NAME": "Alpha"}
    assert mml_service.delete_table("NEW TABLE")
    assert [table["table_name"] for table in mml_service.get_tables_summary()] == ["EXISTING"]


def test_create_table_validates_names_columns_and_duplicates():
    mml_service.import_mml_text("SET EXISTING:ID=1;", "tables.mml")

    for table_name, columns in (("", ["ID"]), ("BAD:TABLE", ["ID"]), ("NEW", []), ("NEW", ["BAD,FIELD"])):
        try:
            mml_service.create_table(table_name, columns)
        except ValueError:
            pass
        else:
            raise AssertionError("invalid table definition should fail")

    try:
        mml_service.create_table("EXISTING", ["ID"])
    except ValueError as exc:
        assert "已存在" in str(exc)
    else:
        raise AssertionError("duplicate table should fail")


def test_manual_table_types_are_persisted_and_values_are_normalized():
    mml_service.create_configuration("Typed Config")
    mml_service.create_table(
        "TYPED",
        ["COUNT", "RATIO", "ACTIVE", "LABEL"],
        {"COUNT": "integer", "RATIO": "decimal", "ACTIVE": "boolean", "LABEL": "string"},
    )

    row_id = mml_service.add_config("TYPED", {"COUNT": "7", "RATIO": "1.5", "ACTIVE": "true", "LABEL": 42})
    summary = mml_service.get_tables_summary()[0]
    row = mml_service.get_config("TYPED", row_id)["config_data"]

    assert summary["column_types"] == {
        "ACTIVE": "boolean",
        "COUNT": "integer",
        "LABEL": "string",
        "RATIO": "decimal",
    }
    assert row == {"COUNT": 7, "RATIO": 1.5, "ACTIVE": True, "LABEL": "42"}

    with pytest.raises(ValueError, match="COUNT"):
        mml_service.add_config("TYPED", {"COUNT": "seven"})


def test_table_properties_can_be_updated_without_losing_mapped_data():
    mml_service.create_configuration("Editable Config")
    mml_service.create_table(
        "OLD TABLE",
        ["ID", "NAME", "REMOVE_ME"],
        {"ID": "string", "NAME": "string", "REMOVE_ME": "string"},
    )
    row_id = mml_service.add_config("OLD TABLE", {"ID": "7", "NAME": "Alpha", "REMOVE_ME": "discarded"})

    updated = mml_service.update_table(
        "OLD TABLE",
        "NEW TABLE",
        ["ID", "DISPLAY_NAME", "ACTIVE"],
        {"ID": "integer", "DISPLAY_NAME": "string", "ACTIVE": "boolean"},
        {"ID": "ID", "DISPLAY_NAME": "NAME"},
    )

    assert updated == {
        "table_name": "NEW TABLE",
        "columns": ["ACTIVE", "DISPLAY_NAME", "ID"],
        "column_types": {"ACTIVE": "boolean", "DISPLAY_NAME": "string", "ID": "integer"},
        "count": 1,
    }
    assert mml_service.get_config("NEW TABLE", row_id)["config_data"] == {
        "ACTIVE": None,
        "DISPLAY_NAME": "Alpha",
        "ID": 7,
    }
    with pytest.raises(ValueError, match="不存在"):
        mml_service.get_config("OLD TABLE", row_id)


def test_incompatible_table_type_update_leaves_definition_unchanged():
    mml_service.create_configuration("Atomic Update")
    mml_service.create_table("ITEMS", ["VALUE"], {"VALUE": "string"})
    mml_service.add_config("ITEMS", {"VALUE": "not-a-number"})

    with pytest.raises(ValueError, match="VALUE"):
        mml_service.update_table("ITEMS", "RENAMED", ["VALUE"], {"VALUE": "integer"})

    assert mml_service.get_tables_summary()[0]["table_name"] == "ITEMS"


def test_selected_rows_can_be_exported_as_csv_and_excel():
    mml_service.import_mml_text(
        'SET CELL:ID=1,NAME="Alpha"; ADD CELL:ID=2,NAME="Beta",POWER=-2.5,CODE="001",FORMULA="=1+1";',
        "cells.mml",
    )
    configs = mml_service.get_configs("CELL", page_size=20)["configs"]
    selected_id = next(row["id"] for row in configs if row["config_data"]["ID"] == 2)

    csv_export = mml_service.export_configurations("csv", "CELL", [selected_id])
    csv_rows = list(csv.DictReader(io.StringIO(csv_export["content"].decode("utf-8-sig"))))
    assert csv_export["filename"].startswith("CELL_selected_")
    assert csv_export["media_type"].startswith("text/csv")
    assert csv_export["count"] == 1
    assert csv_rows == [{"CODE": "001", "FORMULA": "=1+1", "ID": "2", "NAME": "Beta", "POWER": "-2.5"}]

    excel_export = mml_service.export_configurations("xlsx", "CELL", [selected_id])
    workbook = load_workbook(io.BytesIO(excel_export["content"]), read_only=True, data_only=True)
    assert workbook.sheetnames == ["CELL"]
    sheet = workbook["CELL"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, headers["ID"]).value == 2
    assert sheet.cell(2, headers["ID"]).data_type == "n"
    assert sheet.cell(2, headers["POWER"]).value == -2.5
    assert sheet.cell(2, headers["POWER"]).data_type == "n"
    assert sheet.cell(2, headers["CODE"]).value == "001"
    assert sheet.cell(2, headers["CODE"]).data_type == "s"
    assert sheet.cell(2, headers["CODE"]).number_format == "@"
    assert sheet.cell(2, headers["FORMULA"]).value == "=1+1"
    assert sheet.cell(2, headers["FORMULA"]).data_type == "s"
    workbook.close()


def test_all_tables_can_be_exported_as_mml_csv_zip_and_excel():
    mml_service.import_mml_text('SET CELL:ID=1; ADD USER PROFILE:ID=2,NAME="Admin";', "all.mml")

    mml_export = mml_service.export_configurations("mml")
    assert mml_export["count"] == 2
    assert "SET CELL:ID=1;" in mml_export["content"].decode("utf-8-sig")

    csv_export = mml_service.export_configurations("csv")
    assert csv_export["media_type"] == "application/zip"
    with zipfile.ZipFile(io.BytesIO(csv_export["content"])) as archive:
        assert archive.namelist() == ["CELL.csv", "USER PROFILE.csv"]
        assert "ID" in archive.read("CELL.csv").decode("utf-8-sig")

    excel_export = mml_service.export_configurations("xlsx")
    workbook = load_workbook(io.BytesIO(excel_export["content"]), read_only=True, data_only=True)
    assert workbook.sheetnames == ["CELL", "USER PROFILE"]
    assert list(workbook["CELL"].values) == [("ID",), (1,)]
    workbook.close()


def test_export_rejects_empty_selection_and_unknown_format():
    mml_service.import_mml_text("SET CELL:ID=1;", "one.mml")

    for export_format, ids in (("pdf", None), ("csv", [])):
        try:
            mml_service.export_configurations(export_format, "CELL", ids)
        except ValueError:
            pass
        else:
            raise AssertionError("invalid export request should fail")
