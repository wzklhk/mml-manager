# -*- coding: utf-8 -*-
"""
业务逻辑层 (Service)
处理 MML 导入/导出/CRUD 的业务规则。
"""

import csv
import codecs
import io
import json
import os
import re
import zipfile
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from ..mml_parser import parse_mml_stream, parse_mml_text
from ..utils.mml import format_mml_command
from ..utils.tabular import normalize_excel_value, parse_csv_value, write_excel_cell
from .memory_store import store


KEY_FIELD_CANDIDATES = (
    "ID",
    "INDEX",
    "SEQ",
    "SEQUENCE",
    "NAME",
    "MOID",
    "DN",
    "OBJECTID",
)

EXPORT_FORMATS = {"mml", "csv", "xlsx"}


def decode_mml_bytes(content: bytes) -> str:
    """Decode common network-element export encodings."""
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("文件编码无法识别，请使用 UTF-8 或 GB18030 编码")


def _read_mml_text(file_path: str) -> str:
    with open(file_path, "rb") as handle:
        return decode_mml_bytes(handle.read())


def _parse_mml_text(text: str) -> Dict[str, List[Dict]]:
    return parse_mml_text(text)


def _choose_key_fields(before: List[Dict], after: List[Dict]) -> List[str]:
    """选择两侧均存在且能唯一标识记录的字段。"""
    rows = [item["values"] for item in before + after]
    if not rows:
        return []
    common = set(rows[0])
    for row in rows[1:]:
        common.intersection_update(row)
    candidates = [field for field in KEY_FIELD_CANDIDATES if field in common]
    candidates.extend(sorted(field for field in common if field.endswith("_ID") and field not in candidates))
    for field in candidates:
        before_keys = [str(item["values"].get(field, "")) for item in before]
        after_keys = [str(item["values"].get(field, "")) for item in after]
        if (
            all(before_keys + after_keys)
            and len(before_keys) == len(set(before_keys))
            and len(after_keys) == len(set(after_keys))
        ):
            return [field]
    return []


def _record_key(values: Dict, fields: List[str]) -> str:
    return " | ".join(f"{field}={values.get(field, '')}" for field in fields)


def compare_mml_files(baseline_path: str, target_path: str) -> Dict:
    """对比两份 MML 配置，返回表级和字段级差异，不写入数据库。"""
    return compare_mml_texts(_read_mml_text(baseline_path), _read_mml_text(target_path))


def compare_mml_texts(baseline_text: str, target_text: str) -> Dict:
    """Compare two already-loaded MML texts entirely in memory."""
    baseline = _parse_mml_text(baseline_text)
    target = _parse_mml_text(target_text)
    if not baseline and not target:
        raise ValueError("两份文件中都没有找到有效的 SET/ADD 命令")
    return _compare_configurations(baseline, target)


def _compare_configurations(baseline: Dict[str, List[Dict]], target: Dict[str, List[Dict]]) -> Dict:
    table_results = []
    totals = {"added": 0, "removed": 0, "modified": 0, "unchanged": 0}
    for table_name in sorted(set(baseline) | set(target)):
        before = baseline.get(table_name, [])
        after = target.get(table_name, [])
        key_fields = _choose_key_fields(before, after)
        warning = None
        if key_fields:
            before_map = {_record_key(item["values"], key_fields): item["values"] for item in before}
            after_map = {_record_key(item["values"], key_fields): item["values"] for item in after}
        else:
            # 无稳定主键时仍可准确识别完全相同、新增和删除，避免错误地配对为“修改”。
            warning = "未发现唯一标识字段，无法判定字段级修改；差异按完整配置行识别"

            def canonical(item):
                return json.dumps(item["values"], ensure_ascii=False, sort_keys=True)

            before_map = {canonical(item): item["values"] for item in before}
            after_map = {canonical(item): item["values"] for item in after}

        diffs = []
        for key in sorted(set(before_map) | set(after_map)):
            old = before_map.get(key)
            new = after_map.get(key)
            if old is None:
                status, changes = "added", []
            elif new is None:
                status, changes = "removed", []
            elif old != new:
                status = "modified"
                changes = [
                    {"field": field, "before": old.get(field), "after": new.get(field)}
                    for field in sorted(set(old) | set(new))
                    if old.get(field) != new.get(field)
                ]
            else:
                status, changes = "unchanged", []
            totals[status] += 1
            if status != "unchanged":
                diffs.append({"key": key, "status": status, "before": old, "after": new, "changes": changes})

        table_results.append(
            {
                "table_name": table_name,
                "key_fields": key_fields,
                "baseline_count": len(before),
                "target_count": len(after),
                "warning": warning,
                "diffs": diffs,
                "summary": {
                    status: sum(1 for diff in diffs if diff["status"] == status)
                    for status in ("added", "removed", "modified")
                },
            }
        )
    return {"summary": totals, "tables": table_results}


def compare_snapshots(baseline_id: str, target_id: str) -> Dict:
    if not baseline_id or not target_id:
        raise ValueError("请选择两个配置")
    if baseline_id == target_id:
        raise ValueError("请选择两个不同的配置")

    def snapshot_commands(snapshot_id):
        tables, _ = store.snapshot(snapshot_id)
        return {
            table_name: [{"values": row["values"]} for row in table["rows"]] for table_name, table in tables.items()
        }

    return _compare_configurations(snapshot_commands(baseline_id), snapshot_commands(target_id))


def import_mml_file(file_path: str) -> Dict:
    """Parse a file into a persistent configuration set and make it active."""
    with open(file_path, "rb") as handle:
        return import_mml_stream(handle, os.path.basename(file_path))


def import_mml_stream(binary_stream, name: str, network_element: str | None = None) -> Dict:
    """Decode, parse and persist an MML stream with bounded working memory."""
    last_error = None
    for encoding in ("utf-8-sig", "gb18030"):
        binary_stream.seek(0)
        reader = codecs.getreader(encoding)(binary_stream, errors="strict")
        try:
            snapshot, table_names = store.add_snapshot_stream(parse_mml_stream(reader), name, network_element)
            return {
                "message": f"成功解析并持久化 {snapshot['command_count']} 条配置",
                "tables": table_names,
                "total_count": snapshot["command_count"],
                "snapshot": snapshot,
            }
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError("文件编码无法识别，请使用 UTF-8 或 GB18030 编码") from last_error


def _tabular_headers(row, table_name: str) -> Tuple[List[str], List[int]]:
    headers = [str(value).strip() if value is not None else "" for value in row]
    indices = [index for index, header in enumerate(headers) if header]
    selected = [headers[index] for index in indices]
    if len(selected) != len(set(selected)):
        raise ValueError(f"表 {table_name} 包含重复列名")
    return selected, indices


def _iter_csv_commands(binary_stream, name: str):
    binary_stream.seek(0)
    text = decode_mml_bytes(binary_stream.read())
    reader = csv.reader(io.StringIO(text))
    header_row = next(reader, None)
    table_name = os.path.splitext(os.path.basename(name))[0].strip() or "TABLE"
    if header_row is None:
        return
    headers, indices = _tabular_headers(header_row, table_name)
    if not headers:
        return
    for row in reader:
        values = {
            header: parse_csv_value(row[index] if index < len(row) else "") for header, index in zip(headers, indices)
        }
        if any(value not in (None, "") for value in values.values()):
            yield {"cmd_type": "SET", "table": table_name, "values": values}


def _iter_excel_commands(binary_stream):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    binary_stream.seek(0)
    try:
        workbook = load_workbook(binary_stream, read_only=True, data_only=True)
    except (InvalidFileException, OSError, zipfile.BadZipFile) as exc:
        raise ValueError("Excel 文件无效或已损坏") from exc
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            if header_row is None:
                continue
            headers, indices = _tabular_headers(header_row, sheet.title)
            if not headers:
                continue
            for row in rows:
                values = {
                    header: normalize_excel_value(row[index] if index < len(row) else None)
                    for header, index in zip(headers, indices)
                }
                if any(value not in (None, "") for value in values.values()):
                    yield {"cmd_type": "SET", "table": sheet.title, "values": values}
    finally:
        workbook.close()


def import_configuration_stream(binary_stream, name: str, network_element: str | None = None) -> Dict:
    """Import MML, CSV, or XLSX content while preserving scalar value types."""
    extension = os.path.splitext(name)[1].lower()
    if extension in (".mml", ".txt"):
        return import_mml_stream(binary_stream, name, network_element)
    if extension == ".csv":
        commands = _iter_csv_commands(binary_stream, name)
    elif extension == ".xlsx":
        commands = _iter_excel_commands(binary_stream)
    else:
        raise ValueError("只支持 .mml、.txt、.csv 或 .xlsx 格式文件")
    snapshot, table_names = store.add_snapshot_stream(commands, name, network_element)
    return {
        "message": f"成功导入并持久化 {snapshot['command_count']} 条配置",
        "tables": table_names,
        "total_count": snapshot["command_count"],
        "snapshot": snapshot,
    }


def import_mml_text(text: str, name: str = "未命名配置", network_element: str | None = None) -> Dict:
    """Parse text, persist an isolated configuration set, and make it active."""
    tables = _parse_mml_text(text)
    total_count = sum(len(commands) for commands in tables.values())
    if not total_count:
        return {"error": "未找到有效的MML命令"}
    snapshot = store.add_snapshot(tables, name, network_element)
    return {
        "message": f"成功解析并持久化 {total_count} 条配置",
        "tables": list(tables),
        "total_count": total_count,
        "snapshot": snapshot,
    }


def get_snapshots() -> Dict:
    snapshots, active_id = store.list_snapshots()
    return {"snapshots": snapshots, "active_id": active_id}


def create_configuration(name: str) -> Dict:
    if not isinstance(name, str):
        raise ValueError("配置名称格式无效")
    name = name.strip()
    if not name:
        raise ValueError("配置名称不能为空")
    if len(name) > 128 or re.search(r"[\r\n]", name):
        raise ValueError("配置名称不能超过 128 个字符或包含换行")
    snapshot = store.add_snapshot({}, name)
    return {
        "message": "配置创建成功",
        "tables": [],
        "total_count": 0,
        "snapshot": snapshot,
    }


def activate_snapshot(snapshot_id: str) -> Dict:
    return store.activate(snapshot_id)


def delete_snapshot(snapshot_id: str) -> Dict:
    return store.delete_snapshot(snapshot_id)


def get_tables_summary() -> List[Dict]:
    tables, loaded_at = store.table_summaries()
    return [
        {
            "table_name": table["table_name"],
            "columns": table["columns"],
            "column_types": table.get("column_types", {}),
            "count": table["count"],
            "created_at": loaded_at or "",
        }
        for table in tables
    ]


def get_configs(
    table_name: str,
    page: int = 1,
    page_size: int = 20,
    sort_by: str = None,
    sort_order: str = "asc",
    filters: Dict | None = None,
) -> Dict:
    """
    分页获取配置列表。

    Returns:
        {configs, total, page, page_size, total_pages}
    """
    if not table_name:
        return {"configs": [], "total": 0, "page": 1, "page_size": page_size, "total_pages": 1}

    table, loaded_at = store.table_info(table_name)
    columns = table["columns"]
    filters = {
        field: str(value)
        for field, value in (filters or {}).items()
        if field in columns and value is not None and str(value) != ""
    }
    if not sort_by:
        # 默认按网元对象标识排序，避免导入顺序不同导致页面与导出结果难以核对。
        sort_by = next((field for field in KEY_FIELD_CANDIDATES if field in columns), None)
    rows, total = store.query_page(table_name, page, page_size, sort_by, sort_order, filters)
    page = max(1, int(page))
    page_size = max(1, min(int(page_size), store.max_page_size))
    configs = [
        {
            "id": row["id"],
            "table_name": table_name,
            "cmd_type": row["cmd_type"],
            "config_data": {column: row["values"].get(column) for column in columns},
            "created_at": loaded_at or "",
            "updated_at": loaded_at or "",
        }
        for row in rows
    ]
    total_pages = max(1, (total + page_size - 1) // page_size)
    return {
        "configs": configs,
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": total_pages,
    }


def get_config(table_name: str, config_id: int) -> Optional[Dict]:
    _, loaded_at = store.table_info(table_name)
    row = store.get(table_name, config_id)
    return (
        None
        if row is None
        else {
            "id": row["id"],
            "table_name": table_name,
            "cmd_type": row["cmd_type"],
            "config_data": row["values"],
            "created_at": loaded_at or "",
            "updated_at": loaded_at or "",
        }
    )


def _normalize_values_for_types(config_data: Dict, column_types: Dict[str, str]) -> Dict:
    if not isinstance(config_data, dict):
        raise ValueError("配置数据格式无效")
    normalized = dict(config_data)
    for column, data_type in column_types.items():
        value = normalized.get(column)
        if value is None or value == "":
            normalized[column] = None if data_type != "string" else ""
            continue
        try:
            if data_type == "integer":
                if isinstance(value, bool) or float(value) != int(float(value)):
                    raise ValueError
                normalized[column] = int(float(value))
            elif data_type == "decimal":
                normalized[column] = float(value)
            elif data_type == "boolean":
                if isinstance(value, bool):
                    normalized[column] = value
                elif str(value).strip().lower() in {"true", "1", "yes"}:
                    normalized[column] = True
                elif str(value).strip().lower() in {"false", "0", "no"}:
                    normalized[column] = False
                else:
                    raise ValueError
            else:
                normalized[column] = str(value)
        except (TypeError, ValueError):
            raise ValueError(f"字段 {column} 的值不是有效的 {data_type}") from None
    return normalized


def _normalize_typed_values(table_name: str, config_data: Dict) -> Dict:
    table, _ = store.table_info(table_name)
    return _normalize_values_for_types(config_data, table.get("column_types", {}))


def add_config(table_name: str, config_data: Dict) -> int:
    return store.add(table_name, _normalize_typed_values(table_name, config_data))


def update_config(table_name: str, config_id: int, config_data: Dict) -> bool:
    return store.update(table_name, config_id, _normalize_typed_values(table_name, config_data))


def delete_config(table_name: str, config_id: int) -> bool:
    return store.delete(table_name, [config_id]) > 0


def batch_delete_configs(table_name: str, ids: List[int]) -> int:
    return store.delete(table_name, ids)


def _normalize_table_definition(
    table_name: str, columns: List[str], column_types: Dict[str, str] | None
) -> tuple[str, List[str], Dict[str, str]]:
    if not isinstance(table_name, str) or not isinstance(columns, list):
        raise ValueError("表名和字段格式无效")
    table_name = table_name.strip()
    if not table_name:
        raise ValueError("表名不能为空")
    if len(table_name) > 128 or re.search(r"[:;\r\n]", table_name):
        raise ValueError("表名不能超过 128 个字符，且不能包含冒号、分号或换行")
    normalized_columns = sorted({str(column).strip() for column in (columns or []) if str(column).strip()})
    if not normalized_columns:
        raise ValueError("至少需要一个字段")
    if any(len(column) > 128 or re.search(r"[=,;\r\n]", column) for column in normalized_columns):
        raise ValueError("字段名不能超过 128 个字符，且不能包含等号、逗号、分号或换行")
    column_types = column_types or {}
    if not isinstance(column_types, dict):
        raise ValueError("字段类型格式无效")
    supported_types = {"string", "integer", "decimal", "boolean"}
    unknown_columns = set(column_types) - set(normalized_columns)
    if unknown_columns:
        raise ValueError("字段类型包含未定义的字段")
    normalized_types = {column: column_types.get(column, "string") for column in normalized_columns}
    if any(data_type not in supported_types for data_type in normalized_types.values()):
        raise ValueError("字段类型仅支持 string、integer、decimal 或 boolean")
    return table_name, normalized_columns, normalized_types


def create_table(table_name: str, columns: List[str], column_types: Dict[str, str] | None = None) -> Dict:
    table_name, normalized_columns, normalized_types = _normalize_table_definition(table_name, columns, column_types)
    store.create_table(table_name, normalized_columns, normalized_types)
    return {
        "table_name": table_name,
        "columns": normalized_columns,
        "column_types": normalized_types,
        "count": 0,
    }


def update_table(
    original_table_name: str,
    table_name: str,
    columns: List[str],
    column_types: Dict[str, str] | None = None,
    column_mapping: Dict[str, str] | None = None,
) -> Dict:
    if not isinstance(original_table_name, str) or not original_table_name.strip():
        raise ValueError("原表名不能为空")
    original_table_name = original_table_name.strip()
    table_name, normalized_columns, normalized_types = _normalize_table_definition(table_name, columns, column_types)
    current_table, _ = store.table_info(original_table_name)
    existing_columns = set(current_table["columns"])
    if column_mapping is None:
        column_mapping = {column: column for column in normalized_columns}
    if not isinstance(column_mapping, dict):
        raise ValueError("字段映射格式无效")
    if set(column_mapping) - set(normalized_columns) or set(column_mapping.values()) - existing_columns:
        raise ValueError("字段映射包含无效字段")

    tables, _ = store.snapshot()
    source_rows = tables[original_table_name]["rows"]
    normalized_rows = []
    for row in source_rows:
        values = {
            column: row["values"].get(column_mapping[column]) if column in column_mapping else None
            for column in normalized_columns
        }
        try:
            values = _normalize_values_for_types(values, normalized_types)
        except ValueError as exc:
            raise ValueError(f"无法修改表属性：第 {row['id']} 行{exc}") from None
        normalized_rows.append({"id": row["id"], "values": values})

    store.update_table(original_table_name, table_name, normalized_columns, normalized_types, normalized_rows)
    return {
        "table_name": table_name,
        "columns": normalized_columns,
        "column_types": normalized_types,
        "count": current_table["count"],
    }


def delete_table(table_name: str) -> bool:
    if not isinstance(table_name, str):
        raise ValueError("表名格式无效")
    table_name = table_name.strip()
    if not table_name:
        raise ValueError("表名不能为空")
    return store.delete_table(table_name)


def _select_export_tables(table_name: Optional[str] = None, ids: Optional[List[int]] = None) -> Dict:
    """Return the requested tables/rows from the active snapshot."""
    tables, _ = store.snapshot()
    if table_name:
        if table_name not in tables:
            raise ValueError(f"表 {table_name} 不存在")
        tables = {table_name: tables[table_name]}
    elif ids is not None:
        raise ValueError("导出选中配置时必须指定 table_name")

    if ids is not None:
        wanted = set(ids)
        selected_rows = [row for row in tables[table_name]["rows"] if row["id"] in wanted]
        tables[table_name]["rows"] = selected_rows

    tables = {name: table for name, table in tables.items() if table["rows"]}
    if not tables:
        raise ValueError("没有可导出的配置")
    return tables


def _safe_file_stem(name: str) -> str:
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")
    return stem or "config"


def _unique_names(names: List[str], max_length: Optional[int] = None) -> List[str]:
    result = []
    used = set()
    for name in names:
        base = name[:max_length] if max_length else name
        candidate = base
        index = 2
        while candidate.casefold() in used:
            suffix = f"_{index}"
            candidate = f"{base[: max_length - len(suffix)]}{suffix}" if max_length else f"{base}{suffix}"
            index += 1
        used.add(candidate.casefold())
        result.append(candidate)
    return result


def _csv_bytes(table: Dict) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=table["columns"], extrasaction="ignore")
    writer.writeheader()
    for row in table["rows"]:
        writer.writerow({column: row["values"].get(column, "") for column in table["columns"]})
    return output.getvalue().encode("utf-8-sig")


def _export_csv(tables: Dict, selected: bool, timestamp: str) -> Dict:
    if len(tables) == 1:
        table_name, table = next(iter(tables.items()))
        suffix = "_selected" if selected else ""
        return {
            "content": _csv_bytes(table),
            "filename": f"{_safe_file_stem(table_name)}{suffix}_{timestamp}.csv",
            "media_type": "text/csv; charset=utf-8",
        }

    output = io.BytesIO()
    table_names = sorted(tables, key=str.casefold)
    file_names = _unique_names([_safe_file_stem(name) for name in table_names])
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for table_name, file_name in zip(table_names, file_names):
            archive.writestr(f"{file_name}.csv", _csv_bytes(tables[table_name]))
    return {
        "content": output.getvalue(),
        "filename": f"configs_csv_{timestamp}.zip",
        "media_type": "application/zip",
    }


def _export_excel(tables: Dict, selected: bool, timestamp: str) -> Dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    table_names = sorted(tables, key=str.casefold)
    sheet_names = _unique_names(
        [(re.sub(r"[\[\]:*?/\\]", "_", name).strip("'") or "Config") for name in table_names],
        max_length=31,
    )
    header_fill = PatternFill(fill_type="solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")

    for table_name, sheet_name in zip(table_names, sheet_names):
        table = tables[table_name]
        sheet = workbook.create_sheet(sheet_name)
        for column_index, column in enumerate(table["columns"], 1):
            cell = sheet.cell(row=1, column=column_index, value=column)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row_index, row in enumerate(table["rows"], 2):
            for column_index, column in enumerate(table["columns"], 1):
                write_excel_cell(sheet, row_index, column_index, row["values"].get(column))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_index, column in enumerate(table["columns"], 1):
            values = [column] + [str(row["values"].get(column, "") or "") for row in table["rows"][:100]]
            sheet.column_dimensions[get_column_letter(column_index)].width = min(max(map(len, values)) + 2, 50)

    output = io.BytesIO()
    workbook.save(output)
    suffix = "_selected" if selected else ""
    return {
        "content": output.getvalue(),
        "filename": f"configs{suffix}_{timestamp}.xlsx",
        "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }


def export_configurations(
    export_format: str, table_name: Optional[str] = None, ids: Optional[List[int]] = None
) -> Dict:
    """Export selected rows, one table, or the full active snapshot."""
    export_format = (export_format or "").lower()
    if export_format not in EXPORT_FORMATS:
        raise ValueError("导出格式只支持 mml、csv 或 xlsx")
    if ids is not None and not ids:
        raise ValueError("没有选中要导出的配置")

    tables = _select_export_tables(table_name, ids)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if export_format == "mml":
        result = export_selected_rows(table_name, ids) if ids is not None else export_mml(table_name)
        return {
            "content": result["content"].encode("utf-8-sig"),
            "filename": result["filename"],
            "media_type": "text/plain; charset=utf-8",
            "count": sum(len(table["rows"]) for table in tables.values()),
        }
    result = (
        _export_csv(tables, ids is not None, timestamp)
        if export_format == "csv"
        else _export_excel(tables, ids is not None, timestamp)
    )
    result["count"] = sum(len(table["rows"]) for table in tables.values())
    return result


def export_selected_rows(table_name: str, ids: List[int]) -> Dict:
    tables, _ = store.snapshot()
    if table_name not in tables:
        raise ValueError(f"表 {table_name} 不存在")
    wanted = set(ids)
    rows = [row for row in tables[table_name]["rows"] if row["id"] in wanted]
    if not rows:
        raise ValueError("没有可导出的数据")

    lines = []
    lines.append(f"-- 由MML Manager导出（选中行）\n")
    lines.append(f"-- 生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")

    for row in rows:
        line = format_mml_command(row["cmd_type"], table_name, row["values"])
        lines.append(line + "\n")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return {
        "content": "".join(lines),
        "filename": f"{table_name}_selected_{timestamp}.mml",
    }


def export_mml(table_name: Optional[str] = None) -> Dict:
    """
    导出为 MML 格式。

    Returns:
        {content, filename}
    """
    tables, _ = store.snapshot()
    if table_name:
        if table_name not in tables:
            raise ValueError(f"表 {table_name} 不存在")
        tables = {table_name: tables[table_name]}
    if not tables:
        raise ValueError("没有可导出的配置")

    lines = []
    lines.append(f"-- 由MML Manager导出\n")
    lines.append(f"-- 生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")

    for tname, table in tables.items():
        rows = table["rows"]
        if not rows:
            continue

        lines.append(f"-- ===== {tname} =====\n")
        for row in rows:
            line = format_mml_command(row["cmd_type"], tname, row["values"])
            lines.append(line + "\n")
        lines.append("\n")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return {
        "content": "".join(lines),
        "filename": f"export_{timestamp}.mml",
    }
